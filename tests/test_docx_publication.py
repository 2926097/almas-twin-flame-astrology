from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from almas_tfa.docx_publication import (
    DocxPublicationError,
    PERSONAL_PROFILE_ID,
    PROFILE_ID,
    build_authored_report_docx,
    build_personal_authored_report_docx,
    build_report_docx,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX_AVAILABLE = importlib.util.find_spec("docx") is not None


@unittest.skipUnless(DOCX_AVAILABLE, "publication-docx extra not installed")
class DocxPublicationTests(unittest.TestCase):
    def authored(self):
        return json.loads(
            (ROOT / "examples/authored-report.synthetic.json").read_text(
                encoding="utf-8"
            )
        )

    def test_b5_docx_preserves_all_authored_narratives_and_bibliography(self):
        from docx import Document
        from docx.shared import Mm

        authored = self.authored()
        with TemporaryDirectory() as temp:
            output = Path(temp) / "report.docx"
            receipt = build_authored_report_docx(authored, output)

            self.assertTrue(output.is_file())
            self.assertGreater(output.stat().st_size, 0)
            self.assertEqual(receipt.profile_id, PROFILE_ID)
            self.assertEqual(receipt.section_count, 11)
            self.assertEqual(
                receipt.bibliography_count,
                len(authored["bibliography"]),
            )
            self.assertEqual(
                receipt.canonical_fingerprint,
                authored["canonical_fingerprint"],
            )
            self.assertEqual(len(receipt.docx_sha256), 64)

            document = Document(str(output))
            section = document.sections[0]
            self.assertLessEqual(
                abs(section.page_width - Mm(176)),
                500,
            )
            self.assertLessEqual(
                abs(section.page_height - Mm(250)),
                500,
            )

            paragraph_texts = [p.text for p in document.paragraphs]
            for authored_section in authored["sections"]:
                self.assertIn(authored_section["title"], paragraph_texts)
                self.assertIn(authored_section["narrative"], paragraph_texts)

            for entry in authored["bibliography"]:
                self.assertTrue(
                    any(
                        entry["citation_label"] in text
                        for text in paragraph_texts
                    )
                )

            self.assertIn(
                authored["canonical_fingerprint"],
                document.core_properties.comments,
            )
            footer_text = document.sections[0].footer.paragraphs[0].text
            self.assertIn(
                authored["canonical_fingerprint"][:12],
                footer_text,
            )

    def personal_authored(self):
        section_ids = (
            "P01_SYNTHESIS",
            "P02_DATA_METHOD",
            "P03_NATAL_ARCHITECTURE",
            "P09_TEMPORAL",
            "P10_COUNTEREVIDENCE",
            "P11_SOURCES_ATLAS",
        )
        sections = [
            {
                "section_id": section_id,
                "title": f"Sección personal {index}",
                "authoring_state": "AUTHORED",
                "narrative": f"Narrativa personal sintética {index}.",
                "canonical_paths_used": ["natal"],
                "epistemic_classes_used": ["A_CALCULATED"],
                "doctrinal_claim_refs": [],
                "source_refs": [],
                "limitations": [],
            }
            for index, section_id in enumerate(section_ids, start=1)
        ]
        return {
            "schema_version": "1.0.0",
            "document_kind": "ALMAS_PERSONAL_AUTHORED_REPORT",
            "language": "es",
            "report_profile": "EXECUTIVE_PERSONAL_REPORT",
            "canonical_fingerprint": "c" * 64,
            "report_state": "READY",
            "interpretive_center": (
                "PERSONAL_ASTROLOGY_AND_SOURCE_BASED_HERMENEUTICS"
            ),
            "technical_role": (
                "CALCULATION_TRACEABILITY_AND_QUALITY_CONTROL"
            ),
            "metaphysical_scientific_validation_claimed": False,
            "personal_data_minimized": True,
            "canonical_values_mutated": False,
            "new_calculations_performed": False,
            "new_scores_created": False,
            "sections": sections,
            "bibliography": [],
        }

    def test_personal_stars_narrative_is_published_in_b5_without_mutation(self):
        from copy import deepcopy
        from docx import Document
        from test_personal_fixed_stars import Backend
        from test_personal_request_pipeline import request
        from test_personal_authored_report import authored
        from almas_tfa.personal_request_pipeline import build_personal_report_context_from_request
        from almas_tfa.personal_fixed_stars import render_personal_fixed_stars
        from almas_tfa.personal_authored_report import validate_personal_authored_report_trace
        payload = request(); payload["fixed_stars"] = {"enabled": True}
        context = build_personal_report_context_from_request(payload, Backend())
        canonical = context["personal_canonical_analysis"]
        model = context["personal_report_document_model"]
        before = deepcopy(canonical)
        report = authored(canonical, model); report["bibliography"] = []
        layer = canonical["secondary_layers"]["fixed_stars"]
        section = next(s for s in report["sections"] if s["section_id"] == "P08_COMPLEMENTARY_LAYERS")
        section.update(narrative=render_personal_fixed_stars(layer), canonical_paths_used=["secondary_layers"],
                       epistemic_classes_used=["A_CALCULATED", "B_TECHNIQUE"], source_refs=layer["method_source_ids"])
        for source_id in layer["method_source_ids"]:
            report["bibliography"].append(dict(source_id=source_id, citation_label=source_id,
                source_scope="METHOD", author=None, work=None, tradition=None, locator=None, url=None, verification_anchor=None))
        validate_personal_authored_report_trace(report, model, canonical)
        with TemporaryDirectory() as temp:
            target = Path(temp) / "stars.docx"
            receipt = build_personal_authored_report_docx(report, target)
            texts = "\n".join(p.text for p in Document(str(target)).paragraphs)
            self.assertIn("Brady", texts)
            self.assertIn("SUPPORT_ONLY", texts)
            self.assertEqual(receipt.profile_id, PERSONAL_PROFILE_ID)
        self.assertEqual(canonical, before)

    def test_personal_docx_reuses_b5_renderer(self):
        from docx import Document
        from docx.shared import Mm

        authored = self.personal_authored()
        with TemporaryDirectory() as temp:
            output = Path(temp) / "personal.docx"
            receipt = build_personal_authored_report_docx(
                authored,
                output,
            )
            self.assertEqual(
                receipt.profile_id,
                PERSONAL_PROFILE_ID,
            )
            self.assertEqual(receipt.section_count, 6)

            document = Document(str(output))
            section = document.sections[0]
            self.assertLessEqual(
                abs(section.page_width - Mm(176)),
                500,
            )
            self.assertLessEqual(
                abs(section.page_height - Mm(250)),
                500,
            )
            texts = [p.text for p in document.paragraphs]
            for item in authored["sections"]:
                self.assertIn(item["title"], texts)
                self.assertIn(item["narrative"], texts)

    def test_docx_dispatch_selects_personal_profile(self):
        authored = self.personal_authored()
        with TemporaryDirectory() as temp:
            receipt = build_report_docx(
                authored,
                Path(temp) / "personal.docx",
            )
        self.assertEqual(receipt.profile_id, PERSONAL_PROFILE_ID)

    def test_multparagraph_narrative_is_materialized_as_real_paragraphs(self):
        from docx import Document

        authored = self.authored()
        authored["sections"][0]["narrative"] = (
            "Primer párrafo interpretativo.\n\n"
            "Segundo párrafo interpretativo."
        )
        with TemporaryDirectory() as temp:
            output = Path(temp) / "report.docx"
            build_authored_report_docx(authored, output)
            document = Document(str(output))
            texts = [p.text for p in document.paragraphs]

        self.assertIn("Primer párrafo interpretativo.", texts)
        self.assertIn("Segundo párrafo interpretativo.", texts)

    def test_docx_publication_does_not_mutate_authored_report(self):
        authored = self.authored()
        snapshot = json.dumps(
            authored,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        with TemporaryDirectory() as temp:
            build_authored_report_docx(
                authored,
                Path(temp) / "report.docx",
            )
        self.assertEqual(
            json.dumps(
                authored,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ),
            snapshot,
        )

    def test_invalid_document_kind_is_rejected_before_materialization(self):
        authored = self.authored()
        authored["document_kind"] = "OTHER"
        with TemporaryDirectory() as temp:
            with self.assertRaisesRegex(
                DocxPublicationError,
                "document_kind",
            ):
                build_authored_report_docx(
                    authored,
                    Path(temp) / "report.docx",
                )


if __name__ == "__main__":
    unittest.main()
