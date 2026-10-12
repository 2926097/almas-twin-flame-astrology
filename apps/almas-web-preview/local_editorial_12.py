"""Offline M31 12-section editorial workbench, DOCX/PDF publication.

Additive experimental *local* format; does not modify the official eleven-section
authored-report schema/CLI. Authoring requires provenance and human-authored text.
No cloud processing, no automatic doctrinal claims or invented calculations.
"""
from __future__ import annotations
import argparse
from hashlib import sha256
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Mapping

SECTION_ORDER = tuple(f"S{i:02d}_{n}" for i, n in enumerate((
    "SYNTHESIS", "DATA_METHOD", "NUMERIC_ONTOLOGY", "STRUCTURE",
    "RELATIONAL", "DIFFERENTIAL", "TEMPORAL", "ROBUSTNESS",
    "DOCTRINE", "FINAL_SYNTHESIS", "SOURCES_APPENDICES", "PHASE_DYNAMICS"), start=1))
FORMAT = "ALMAS_LOCAL_EDITORIAL_12_V1"
MAX_SECTION_CHARS = 50000
ALLOWED_CLASSES = {
    "A_CALCULATED", "B_TECHNIQUE", "C_DOCTRINE", "D_CONTEMPORARY_USAGE",
    "E_PROJECT_HYPOTHESIS",
}

class EditorialGateError(ValueError):
    pass

def digest(value: Mapping[str, Any]) -> str:
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
        separators=(",", ":"), default=str).encode("utf-8")).hexdigest()

def verify_m31(canonical: Mapping[str, Any], gate: Mapping[str, Any],
               model: Mapping[str, Any]) -> str:
    fp = digest(canonical)
    if (gate.get("state") not in ("READY", "PARTIAL") or
        gate.get("reportable") is not True or
        gate.get("canonical_values_mutated") is not False or
        gate.get("canonical_fingerprint") != fp or
        model.get("report_state") != gate.get("state") or
        model.get("canonical_fingerprint") != fp or
        model.get("canonical_fingerprint_verified") is not True or
        model.get("canonical_values_mutated") is not False or
        model.get("canonical_values_embedded") is not False or
        model.get("prose_generated") is not False):
        raise EditorialGateError("M30_M31_INTEGRITY_INVALID")
    sections = model.get("sections")
    if (not isinstance(sections, list) or len(sections) != 12 or
        tuple(item.get("section_id") for item in sections if isinstance(item, dict)) != SECTION_ORDER or
        any(item.get("order") != i + 1 for i, item in enumerate(sections))):
        raise EditorialGateError("M31_TWELVE_SECTIONS_INVALID")
    return fp

def make_scaffold(canonical: Mapping[str, Any],gate: Mapping[str, Any],
                  model: Mapping[str, Any]) -> dict[str, Any]:
    fp = verify_m31(canonical, gate, model)
    return {
        "format": FORMAT,
        "canonical_fingerprint": fp,
        "report_state": gate["state"],
        "language": "es",
        "metaphysical_scientific_validation_claimed": False,
        "new_calculations_performed": False,
        "new_scores_created": False,
        "bibliography": [],
        "sections": [
            {
                "section_id": s["section_id"], "title": s["title"],
                "narrative": "", "canonical_paths_used": [],
                "epistemic_classes_used": [], "source_refs": [],
                "limitations": ["Borrador pendiente de autoría con evidencia trazable."]
            }
            for s in model["sections"]
        ],
    }

def validate_authored(canonical: Mapping[str, Any], gate: Mapping[str, Any],
                      model: Mapping[str, Any], authored: Mapping[str, Any]) -> None:
    fp = verify_m31(canonical, gate, model)
    if (authored.get("format") != FORMAT or
        authored.get("canonical_fingerprint") != fp or
        authored.get("report_state") != gate["state"] or
        authored.get("language") != "es" or
        authored.get("metaphysical_scientific_validation_claimed") is not False or
        authored.get("new_calculations_performed") is not False or
        authored.get("new_scores_created") is not False):
        raise EditorialGateError("EDITORIAL_HEADER_INVALID")
    sections = authored.get("sections")
    if (not isinstance(sections, list) or len(sections) != 12 or
        tuple(s.get("section_id") for s in sections if isinstance(s,dict)) != SECTION_ORDER):
        raise EditorialGateError("EDITORIAL_SECTIONS_INVALID")
    bibliography = authored.get("bibliography")
    if not isinstance(bibliography, list):
        raise EditorialGateError("EDITORIAL_BIBLIOGRAPHY_INVALID")
    source_ids = []
    for item in bibliography:
        if not isinstance(item, dict) or not isinstance(item.get("source_id"), str) or not item["source_id"]:
            raise EditorialGateError("EDITORIAL_SOURCE_INVALID")
        if not isinstance(item.get("citation_label"),str) or not item["citation_label"]:
            raise EditorialGateError("EDITORIAL_SOURCE_INVALID")
        source_ids.append(item["source_id"])
    if len(source_ids) != len(set(source_ids)):
        raise EditorialGateError("EDITORIAL_DUPLICATE_SOURCE")
    for spec, chapter in zip(model["sections"], sections):
        if not isinstance(chapter, dict):
            raise EditorialGateError("EDITORIAL_SECTION_INVALID")
        if chapter.get("title") != spec.get("title"):
            raise EditorialGateError("EDITORIAL_SECTION_TITLE_MISMATCH")
        narrative=chapter.get("narrative")
        paths=chapter.get("canonical_paths_used")
        classes=chapter.get("epistemic_classes_used")
        references=chapter.get("source_refs")
        limitations=chapter.get("limitations")
        if (not isinstance(narrative,str) or len(narrative)>MAX_SECTION_CHARS or
            not isinstance(paths,list) or len(paths)!=len(set(paths)) or
            not all(isinstance(v,str) and v in spec.get("available_paths",[]) for v in paths) or
            not isinstance(classes,list) or len(classes)!=len(set(classes)) or
            not all(isinstance(v,str) and v in ALLOWED_CLASSES and v in spec.get("epistemic_classes_allowed",[]) for v in classes) or
            not isinstance(references,list) or len(references)!=len(set(references)) or
            not all(isinstance(v,str) and v in source_ids for v in references) or
            not isinstance(limitations,list) or
            not all(isinstance(v,str) and v for v in limitations)):
            raise EditorialGateError("EDITORIAL_SECTION_PROVENANCE_INVALID")
        if spec.get("section_state")=="NOT_AVAILABLE":
            if narrative or paths or classes or references:
                raise EditorialGateError("EDITORIAL_UNAVAILABLE_CANNOT_BE_AUTHORED")
        elif not narrative.strip() or not paths or not classes:
            raise EditorialGateError("EDITORIAL_AVAILABLE_REQUIRES_TEXT_AND_REFERENCES")
        if set(classes) & {"C_DOCTRINE","D_CONTEMPORARY_USAGE"} and not references:
            raise EditorialGateError("EDITORIAL_DOCTRINAL_SOURCE_REQUIRED")

def render_docx(authored: Mapping[str, Any], output: Path) -> str:
    try:
        from docx import Document
        from docx.shared import Mm, Pt
    except ImportError as exc:
        raise EditorialGateError("PYTHON_DOCX_REQUIRED") from exc
    doc=Document()
    section=doc.sections[0]
    section.page_width=Mm(176);section.page_height=Mm(250)
    section.left_margin=Mm(18);section.right_margin=Mm(18)
    section.top_margin=Mm(19);section.bottom_margin=Mm(19)
    normal=doc.styles["Normal"];normal.font.name="Liberation Serif";normal.font.size=Pt(10.5)
    doc.add_heading("ALMAS · Informe editorial local",level=0)
    doc.add_paragraph("Documento interpretativo redactado y verificado contra M30/M31. "
        "No demuestra científicamente hipótesis metafísicas.")
    doc.add_paragraph("Estado: "+authored["report_state"]+
        " · Huella canónica: "+authored["canonical_fingerprint"])
    for chapter in authored["sections"]:
        doc.add_heading(chapter["title"],level=1)
        if chapter["narrative"]:
            for paragraph in chapter["narrative"].split("\n"):
                if paragraph.strip():doc.add_paragraph(paragraph)
        else:
            doc.add_paragraph("Sección no disponible en el canónico. No se redacta interpretación.")
        doc.add_paragraph("Rutas utilizadas: "+
            (", ".join(chapter["canonical_paths_used"]) or "Ninguna"))
        doc.add_paragraph("Clases epistemológicas: "+
            (", ".join(chapter["epistemic_classes_used"]) or "Ninguna"))
        if chapter["source_refs"]:
            doc.add_paragraph("Fuentes: "+", ".join(chapter["source_refs"]))
        if chapter["limitations"]:
            doc.add_paragraph("Limitaciones: "+"; ".join(chapter["limitations"]))
    doc.add_heading("Fuentes editoriales",level=1)
    for source in authored["bibliography"]:
        doc.add_paragraph(source["source_id"]+": "+source["citation_label"])
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    return sha256(output.read_bytes()).hexdigest()

def render_pdf_from_docx(docx_path: Path, pdf_path: Path, *,
                         fingerprint: str, section_titles: list[str]) -> dict[str, Any]:
    from almas_tfa.pdf_publication import convert_docx_to_pdf, preflight_pdf
    convert_docx_to_pdf(docx_path,pdf_path)
    result=preflight_pdf(pdf_path,required_text=[fingerprint,*section_titles])
    if not result.preflight_passed:
        pdf_path.unlink(missing_ok=True)
        raise EditorialGateError("PDF_B5_PREFLIGHT_FAILED")
    return {"pdf_sha256": sha256(pdf_path.read_bytes()).hexdigest(),
        "pages": result.page_count, "preflight_passed":True}

def _read_json(path: Path) -> dict[str, Any]:
    obj=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(obj,dict):
        raise EditorialGateError("INPUT_OBJECT_REQUIRED")
    return obj

def main() -> int:
    parser=argparse.ArgumentParser(description="ALMAS offline twelve-section authoring; no cloud API")
    parser.add_argument("canonical",type=Path)
    parser.add_argument("m30",type=Path)
    parser.add_argument("m31",type=Path)
    parser.add_argument("--scaffold-out",type=Path)
    parser.add_argument("--authored",type=Path)
    parser.add_argument("--docx-out",type=Path)
    parser.add_argument("--pdf-out",type=Path)
    args=parser.parse_args()
    if bool(args.scaffold_out)==bool(args.authored):
        parser.error("Select exactly one of --scaffold-out or --authored")
    canonical,gate,model=map(_read_json,(args.canonical,args.m30,args.m31))
    # Schema validation is authoritative for each official document.
    from schema_gate import build_validator,validate_document
    root=Path(__file__).resolve().parents[2]
    for filename,obj in (("canonical-analysis.schema.json",canonical),
        ("report-gate-output.schema.json",gate),
        ("report-document-model.schema.json",model)):
        if validate_document(build_validator(root,filename),obj):
            raise EditorialGateError("OFFICIAL_INPUT_SCHEMA_INVALID")
    if args.scaffold_out:
        result=make_scaffold(canonical,gate,model)
        args.scaffold_out.parent.mkdir(parents=True,exist_ok=True)
        args.scaffold_out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print("ALMAS editorial scaffold: CREATED, 12 sections, not authored")
        return 0
    if not args.docx_out and not args.pdf_out:
        parser.error("Provide --docx-out or --pdf-out")
    authored=_read_json(args.authored)
    validate_authored(canonical,gate,model,authored)
    with TemporaryDirectory(prefix="almas-local-editorial-") as tmp:
        docx_path=args.docx_out if args.docx_out else Path(tmp)/"report.docx"
        digest_docx=render_docx(authored,docx_path)
        if args.pdf_out:
            render_pdf_from_docx(docx_path,args.pdf_out,
                fingerprint=authored["canonical_fingerprint"],
                section_titles=[s["title"] for s in authored["sections"]])
    print("ALMAS local edited DOCX/PDF: VALIDATED (no remote publication)")
    return 0

if __name__=="__main__":
    raise SystemExit(main())
