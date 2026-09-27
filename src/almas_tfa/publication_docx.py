from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from .publication_authoring import AUTHORING_PRINCIPLE, SECTION_IDS


ROLE_LABELS = {
    "ASTROLOGICAL_READING": "Lectura astrológica",
    "METAPHYSICAL_HERMENEUTICS": "Hermenéutica metafísica",
    "DOCTRINAL_INTERPRETATION": "Lectura doctrinal",
    "TECHNICAL_SUPPORT": "Soporte técnico",
    "LIMITATION": "Límite interpretativo",
    "SOURCE_NOTE": "Nota de fuentes",
}


def _mm(value: float):
    from docx.shared import Mm
    return Mm(value)


def _configure_styles(document) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 11)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)

    document.styles["Title"].paragraph_format.space_after = Pt(18)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 2"].paragraph_format.space_before = Pt(7)
    document.styles["Heading 2"].paragraph_format.space_after = Pt(3)

    for name in ("ALMAS Technical Support", "ALMAS Source Note", "ALMAS Limitation"):
        if name not in document.styles:
            style = document.styles.add_style(name, 1)
        else:
            style = document.styles[name]
        style.base_style = document.styles["Normal"]
        style.font.size = Pt(9)
        style.font.italic = True
        style.paragraph_format.left_indent = _mm(5)
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)

    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _block_style(role: str) -> str:
    if role == "TECHNICAL_SUPPORT":
        return "ALMAS Technical Support"
    if role == "SOURCE_NOTE":
        return "ALMAS Source Note"
    if role == "LIMITATION":
        return "ALMAS Limitation"
    return "Normal"


def render_publication_docx(
    manuscript: Mapping[str, Any],
    output_path: str | Path,
    *,
    profile: str = "B5_BOOK",
) -> Path:
    if manuscript.get("authoring_principle") != AUTHORING_PRINCIPLE:
        raise ValueError("Principio editorial ALMAS desconocido.")
    if [s.get("section_id") for s in manuscript.get("sections", [])] != list(SECTION_IDS):
        raise ValueError("El manuscrito no conserva las 11 secciones canónicas.")

    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    section = document.sections[0]
    if profile == "B5_BOOK":
        section.page_width = _mm(176)
        section.page_height = _mm(250)
        section.top_margin = _mm(18)
        section.bottom_margin = _mm(18)
        section.left_margin = _mm(20)
        section.right_margin = _mm(18)
    elif profile == "A4_REPORT":
        section.page_width = _mm(210)
        section.page_height = _mm(297)
        section.top_margin = _mm(20)
        section.bottom_margin = _mm(20)
        section.left_margin = _mm(22)
        section.right_margin = _mm(20)
    else:
        raise ValueError(f"Perfil de publicación desconocido: {profile}")

    _configure_styles(document)

    title = document.add_paragraph(style="Title")
    title.add_run(str(manuscript["title"]))
    subtitle = manuscript.get("subtitle")
    if subtitle:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(subtitle))
        run.italic = True
        run.font.size = Pt(11)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Astrología y hermenéutica metafísica con trazabilidad de fuentes")
    run.italic = True
    run.font.size = Pt(9)

    document.add_page_break()

    for index, section_data in enumerate(manuscript["sections"]):
        document.add_heading(str(section_data["title"]), level=1)
        if section_data.get("section_state") != "READY":
            state = document.add_paragraph(style="ALMAS Limitation")
            state.add_run(
                f"Estado de sección: {section_data.get('section_state')}."
            )

        for block in section_data.get("blocks", []):
            role = str(block["role"])
            document.add_heading(ROLE_LABELS.get(role, role), level=2)
            paragraph = document.add_paragraph(
                str(block["text"]),
                style=_block_style(role),
            )

            traces: list[str] = []
            if block.get("canonical_paths"):
                traces.append(
                    "Rutas: " + ", ".join(str(x) for x in block["canonical_paths"])
                )
            if block.get("source_ids"):
                traces.append(
                    "Fuentes: " + ", ".join(str(x) for x in block["source_ids"])
                )
            if traces:
                note = document.add_paragraph(style="ALMAS Source Note")
                note.add_run(" · ".join(traces))

        if index < len(manuscript["sections"]) - 1:
            document.add_page_break()

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output
