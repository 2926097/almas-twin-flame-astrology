from __future__ import annotations

from dataclasses import asdict, dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any, Mapping, Sequence

from .personal_reporting import PROFILE_SECTIONS as PERSONAL_PROFILE_SECTIONS


PROFILE_ID = "ALMAS_B5_BOOK_V1"
PERSONAL_PROFILE_ID = "ALMAS_B5_PERSONAL_BOOK_V1"
SECTION_ORDER = (
    "S01_SYNTHESIS",
    "S02_DATA_METHOD",
    "S03_NUMERIC_ONTOLOGY",
    "S04_STRUCTURE",
    "S05_RELATIONAL",
    "S06_DIFFERENTIAL",
    "S07_TEMPORAL",
    "S08_ROBUSTNESS",
    "S09_DOCTRINE",
    "S10_FINAL_SYNTHESIS",
    "S11_SOURCES_APPENDICES",
)


class DocxPublicationError(ValueError):
    pass


@dataclass(frozen=True)
class DocxPublicationReceipt:
    profile_id: str
    output_path: str
    canonical_fingerprint: str
    section_count: int
    bibliography_count: int
    docx_sha256: str

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def _require_docx():
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError as exc:
        raise RuntimeError(
            "La publicación DOCX requiere el extra opcional "
            "'publication-docx' (python-docx==1.2.0)."
        ) from exc
    return (
        Document,
        WD_STYLE_TYPE,
        WD_ALIGN_PARAGRAPH,
        OxmlElement,
        qn,
        Mm,
        Pt,
    )


def _validate_common_surface(authored_report: Mapping[str, Any]) -> None:
    if authored_report.get("language") != "es":
        raise DocxPublicationError(
            "La capa DOCX sólo publica informes redactados en español."
        )
    if authored_report.get("canonical_values_mutated") is not False:
        raise DocxPublicationError(
            "No se publica un informe que mutó el canonical."
        )
    if authored_report.get("new_calculations_performed") is not False:
        raise DocxPublicationError(
            "La publicación no admite cálculos posteriores a la autoría."
        )
    if authored_report.get("new_scores_created") is not False:
        raise DocxPublicationError(
            "La publicación no admite scores creados en autoría."
        )

    fingerprint = authored_report.get("canonical_fingerprint")
    if not isinstance(fingerprint, str) or len(fingerprint) != 64:
        raise DocxPublicationError("canonical_fingerprint inválido.")

    sections = authored_report.get("sections")
    if not isinstance(sections, Sequence) or isinstance(
        sections, (str, bytes)
    ):
        raise DocxPublicationError("sections debe ser una secuencia.")


def _section_ids(authored_report: Mapping[str, Any]) -> tuple[Any, ...]:
    return tuple(
        section.get("section_id")
        for section in authored_report.get("sections", [])
        if isinstance(section, Mapping)
    )


def _validate_relational_surface(
    authored_report: Mapping[str, Any],
) -> None:
    _validate_common_surface(authored_report)
    if authored_report.get("document_kind") != "ALMAS_AUTHORED_REPORT":
        raise DocxPublicationError(
            "document_kind no es ALMAS_AUTHORED_REPORT."
        )
    if _section_ids(authored_report) != SECTION_ORDER:
        raise DocxPublicationError(
            "El DOCX relacional debe conservar el orden exacto "
            "de las once secciones."
        )


def _validate_personal_surface(
    authored_report: Mapping[str, Any],
) -> None:
    _validate_common_surface(authored_report)
    if (
        authored_report.get("document_kind")
        != "ALMAS_PERSONAL_AUTHORED_REPORT"
    ):
        raise DocxPublicationError(
            "document_kind no es ALMAS_PERSONAL_AUTHORED_REPORT."
        )
    if authored_report.get("personal_data_minimized") is not True:
        raise DocxPublicationError(
            "El informe personal debe conservar personal_data_minimized=true."
        )
    profile = authored_report.get("report_profile")
    expected = PERSONAL_PROFILE_SECTIONS.get(str(profile))
    if expected is None:
        raise DocxPublicationError(
            f"Perfil personal desconocido: {profile}"
        )
    if _section_ids(authored_report) != expected:
        raise DocxPublicationError(
            "El DOCX personal debe conservar el orden exacto "
            "de su report_profile."
        )


def _surface_profile(
    authored_report: Mapping[str, Any],
) -> tuple[str, str, str]:
    kind = authored_report.get("document_kind")
    if kind == "ALMAS_AUTHORED_REPORT":
        _validate_relational_surface(authored_report)
        return (
            PROFILE_ID,
            "ALMAS · Informe interpretativo",
            "Astrología relacional y hermenéutica metafísica basada en fuentes",
        )
    if kind == "ALMAS_PERSONAL_AUTHORED_REPORT":
        _validate_personal_surface(authored_report)
        return (
            PERSONAL_PROFILE_ID,
            "ALMAS · Informe astrológico personal",
            "Astrología natal y hermenéutica basada en fuentes",
        )
    raise DocxPublicationError(
        f"document_kind no publicable: {kind}"
    )
def _add_page_field(paragraph: Any, OxmlElement: Any, qn: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def _set_styles(document: Any, WD_STYLE_TYPE: Any, Pt: Any) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Liberation Sans"
    title.font.size = Pt(24)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Liberation Sans"
    subtitle.font.size = Pt(12)
    subtitle.paragraph_format.space_after = Pt(12)

    heading = styles["Heading 1"]
    heading.font.name = "Liberation Sans"
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True

    heading2 = styles["Heading 2"]
    heading2.font.name = "Liberation Sans"
    heading2.font.size = Pt(11)
    heading2.font.bold = True
    heading2.paragraph_format.space_before = Pt(8)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True

    trace = styles.add_style("ALMAS Trace", WD_STYLE_TYPE.PARAGRAPH)
    trace.font.name = "Liberation Sans"
    trace.font.size = Pt(8)
    trace.font.italic = True
    trace.paragraph_format.space_before = Pt(5)
    trace.paragraph_format.space_after = Pt(3)

    bibliography = styles.add_style(
        "ALMAS Bibliography",
        WD_STYLE_TYPE.PARAGRAPH,
    )
    bibliography.font.name = "Liberation Serif"
    bibliography.font.size = Pt(9)
    bibliography.paragraph_format.left_indent = Pt(12)
    bibliography.paragraph_format.first_line_indent = Pt(-12)
    bibliography.paragraph_format.space_after = Pt(5)


def _citation_index(authored_report: Mapping[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    bibliography = authored_report.get("bibliography", [])
    if not isinstance(bibliography, list):
        return result
    for entry in bibliography:
        if not isinstance(entry, Mapping):
            continue
        source_id = entry.get("source_id")
        label = entry.get("citation_label")
        if isinstance(source_id, str) and isinstance(label, str):
            result[source_id] = label
    return result


def _trace_line(section: Mapping[str, Any], citations: Mapping[str, str]) -> str:
    parts: list[str] = []
    evidence = section.get("evidence_refs", [])
    claims = section.get("doctrinal_claim_refs", [])
    sources = section.get("source_refs", [])
    if evidence:
        parts.append("Evidencia: " + ", ".join(str(x) for x in evidence))
    if claims:
        parts.append("Claims: " + ", ".join(str(x) for x in claims))
    if sources:
        labels = [citations.get(str(x), str(x)) for x in sources]
        parts.append("Fuentes: " + "; ".join(labels))
    return " · ".join(parts)


def _bibliography_text(entry: Mapping[str, Any]) -> str:
    label = str(entry.get("citation_label") or entry.get("source_id") or "Fuente")
    locator = entry.get("locator")
    url = entry.get("url")
    pieces = [label]
    if isinstance(locator, str) and locator:
        pieces.append(locator)
    if isinstance(url, str) and url:
        pieces.append(url)
    return ". ".join(pieces)


def _narrative_paragraphs(narrative: str) -> list[str]:
    normalized = narrative.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [
        block.strip()
        for block in normalized.split("\n\n")
        if block.strip()
    ]
    return blocks or [narrative]


def _build_b5_docx(
    authored_report: Mapping[str, Any],
    output_path: str | Path,
    *,
    profile_id: str,
    title: str,
    subtitle: str,
) -> DocxPublicationReceipt:
    (
        Document,
        WD_STYLE_TYPE,
        WD_ALIGN_PARAGRAPH,
        OxmlElement,
        qn,
        Mm,
        Pt,
    ) = _require_docx()

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(176)
    section.page_height = Mm(250)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    _set_styles(document, WD_STYLE_TYPE, Pt)

    fingerprint = str(authored_report["canonical_fingerprint"])
    properties = document.core_properties
    properties.title = title
    properties.subject = subtitle
    properties.keywords = (
        "ALMAS, astrología, hermenéutica metafísica, informe B5"
    )
    properties.comments = (
        f"ALMAS authored_report · canonical_fingerprint={fingerprint} · "
        f"profile={profile_id}"
    )

    p = document.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)

    p = document.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(subtitle)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        f"Estado del informe: {authored_report.get('report_state', 'READY')}"
    )

    p = document.add_paragraph(style="ALMAS Trace")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Huella canónica: {fingerprint}")

    document.add_page_break()

    contents = document.add_paragraph()
    contents.add_run("Contenido").bold = True
    for index, authored_section in enumerate(
        authored_report["sections"],
        start=1,
    ):
        document.add_paragraph(
            f"{index}. "
            f"{authored_section.get('title', authored_section.get('section_id'))}"
        )

    citations = _citation_index(authored_report)

    for authored_section in authored_report["sections"]:
        heading = document.add_paragraph(
            str(
                authored_section.get("title")
                or authored_section["section_id"]
            ),
            style="Heading 1",
        )
        heading.paragraph_format.keep_with_next = True

        if (
            authored_section.get("authoring_state")
            == "OMITTED_NOT_AVAILABLE"
        ):
            p = document.add_paragraph()
            p.add_run(
                "Sección no disponible en el análisis canónico."
            ).italic = True
            continue

        narrative = authored_section.get("narrative")
        if not isinstance(narrative, str) or not narrative:
            raise DocxPublicationError(
                f"{authored_section['section_id']}: narrative vacío."
            )
        for paragraph_text in _narrative_paragraphs(narrative):
            document.add_paragraph(paragraph_text)

        limitations = authored_section.get("limitations", [])
        if limitations:
            document.add_paragraph(
                "Límites de interpretación",
                style="Heading 2",
            )
            for item in limitations:
                document.add_paragraph(
                    str(item),
                    style="List Bullet",
                )

        trace = _trace_line(authored_section, citations)
        if trace:
            document.add_paragraph(
                "Trazabilidad · " + trace,
                style="ALMAS Trace",
            )

    bibliography = authored_report.get("bibliography", [])
    if bibliography:
        document.add_paragraph(
            "Bibliografía utilizada",
            style="Heading 2",
        )
        for entry in bibliography:
            if isinstance(entry, Mapping):
                document.add_paragraph(
                    _bibliography_text(entry),
                    style="ALMAS Bibliography",
                )

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"ALMAS · {fingerprint[:12]} · ")
    _add_page_field(footer, OxmlElement, qn)

    document.save(str(output))
    digest = sha256(output.read_bytes()).hexdigest()
    return DocxPublicationReceipt(
        profile_id=profile_id,
        output_path=str(output),
        canonical_fingerprint=fingerprint,
        section_count=len(authored_report["sections"]),
        bibliography_count=(
            len(bibliography)
            if isinstance(bibliography, list)
            else 0
        ),
        docx_sha256=digest,
    )


def build_authored_report_docx(
    authored_report: Mapping[str, Any],
    output_path: str | Path,
    *,
    title: str = "ALMAS · Informe interpretativo",
    subtitle: str = (
        "Astrología relacional y hermenéutica metafísica basada en fuentes"
    ),
) -> DocxPublicationReceipt:
    """Materializa el authored_report relacional como DOCX B5."""

    _validate_relational_surface(authored_report)
    return _build_b5_docx(
        authored_report,
        output_path,
        profile_id=PROFILE_ID,
        title=title,
        subtitle=subtitle,
    )


def build_personal_authored_report_docx(
    authored_report: Mapping[str, Any],
    output_path: str | Path,
    *,
    title: str = "ALMAS · Informe astrológico personal",
    subtitle: str = "Astrología natal y hermenéutica basada en fuentes",
) -> DocxPublicationReceipt:
    """Materializa el authored_report personal con el mismo motor B5."""

    _validate_personal_surface(authored_report)
    return _build_b5_docx(
        authored_report,
        output_path,
        profile_id=PERSONAL_PROFILE_ID,
        title=title,
        subtitle=subtitle,
    )


def build_report_docx(
    authored_report: Mapping[str, Any],
    output_path: str | Path,
    *,
    title: str | None = None,
    subtitle: str | None = None,
) -> DocxPublicationReceipt:
    """Dispatch de publicación DOCX según document_kind."""

    profile_id, default_title, default_subtitle = _surface_profile(
        authored_report
    )
    return _build_b5_docx(
        authored_report,
        output_path,
        profile_id=profile_id,
        title=title or default_title,
        subtitle=subtitle or default_subtitle,
    )
